from pathlib import Path
import math
import traceback
import unicodedata

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt


EMU_PER_INCH = 914400


class PptAutoFitService:

    
    def autofit_pptx(
        self,
        pptx_path: Path,
        output_path: Path | None = None,
        max_passes: int = 3
    ) -> dict:
        pptx_path = Path(pptx_path)

        if output_path is None:
            output_path = pptx_path
        else:
            output_path = Path(output_path)

        report = {
            "pptx_path": str(pptx_path),
            "output_path": str(output_path),
            "slides": [],
            "warnings": [],
            "status": "ok"
        }

        try:
            prs = Presentation(str(pptx_path))
        except Exception as error:
            report["status"] = "failed"
            report["warnings"].append(
                {
                    "stage": "open",
                    "message": f"Failed to open PPTX for autofit: {str(error)}"
                }
            )
            return report

        slide_width = int(prs.slide_width)
        slide_height = int(prs.slide_height)

        for slide_index, slide in enumerate(prs.slides):
            slide_report = {
                "slide_index": slide_index,
                "adjusted_shapes": 0,
                "overlap_repairs": 0,
                "warnings": []
            }

            try:
                text_shapes = self._collect_text_shapes(slide)

                for item in text_shapes:
                    adjusted = self._autofit_shape(
                        item=item,
                        slide_width=slide_width,
                        slide_height=slide_height
                    )

                    if adjusted:
                        slide_report["adjusted_shapes"] += 1

                for _ in range(max_passes):
                    repaired = self._resolve_overlaps(
                        text_shapes=text_shapes,
                        slide_width=slide_width,
                        slide_height=slide_height
                    )

                    slide_report["overlap_repairs"] += repaired

                    if repaired == 0:
                        break

                # avoid text outside
                for item in text_shapes:
                    self._clamp_shape_to_slide(
                        item=item,
                        slide_width=slide_width,
                        slide_height=slide_height
                    )

            except Exception as error:
                slide_report["warnings"].append(
                    {
                        "message": str(error),
                        "traceback": traceback.format_exc()
                    }
                )

            report["slides"].append(slide_report)

        try:
            prs.save(str(output_path))
        except Exception as error:
            report["status"] = "failed"
            report["warnings"].append(
                {
                    "stage": "save",
                    "message": f"Failed to save autofitted PPTX: {str(error)}"
                }
            )

        return report

    def _collect_text_shapes(self, slide) -> list[dict]:
        result = []

        for shape in self._iter_shapes(slide.shapes):
            if not getattr(shape, "has_text_frame", False):
                continue

            text = self._shape_text(shape)

            if not text:
                continue

            try:
                width = int(shape.width)
                height = int(shape.height)
                left = int(shape.left)
                top = int(shape.top)
            except Exception:
                continue

            if width <= 0 or height <= 0:
                continue

            role = self._guess_shape_role(shape)

            # skip c txt
            if role == "decorative":
                continue

            result.append(
                {
                    "shape": shape,
                    "role": role,
                    "text": text,
                    "left": left,
                    "top": top,
                    "width": width,
                    "height": height
                }
            )

        return result

    def _iter_shapes(self, shapes):
        for shape in shapes:
            yield shape

            if hasattr(shape, "shapes"):
                try:
                    for child in self._iter_shapes(shape.shapes):
                        yield child
                except Exception:
                    pass

    def _shape_text(self, shape) -> str:
        try:
            text = shape.text_frame.text
        except Exception:
            try:
                text = shape.text
            except Exception:
                text = ""

        return str(text).strip()

    def _guess_shape_role(self, shape) -> str:
        try:
            left = int(shape.left)
            top = int(shape.top)
            width = int(shape.width)
            height = int(shape.height)
        except Exception:
            return "body"

        max_font_size = self._max_font_size(shape)
        text = self._shape_text(shape)

        if width > 0 and height / width >= 2.3 and width < Inches(1.2):
            return "decorative"

        if top < Inches(1.5) and max_font_size >= 22:
            return "title"

        if top < Inches(2.2) and max_font_size >= 14 and len(text) <= 120:
            return "subtitle"

        if len(text) <= 25 and height < Inches(0.7):
            return "label"

        return "body"

    def _autofit_shape(
        self,
        item: dict,
        slide_width: int,
        slide_height: int
    ) -> bool:
        shape = item["shape"]
        role = item["role"]

        changed = False

        try:
            frame = shape.text_frame
        except Exception:
            return False

        # give space
        try:
            frame.word_wrap = True
            frame.vertical_anchor = MSO_ANCHOR.TOP
            frame.margin_left = Inches(0.03)
            frame.margin_right = Inches(0.03)
            frame.margin_top = Inches(0.02)
            frame.margin_bottom = Inches(0.02)
            changed = True
        except Exception:
            pass

      
        try:
            for paragraph in frame.paragraphs:
                paragraph.space_after = Pt(1)
                paragraph.space_before = Pt(0)

                try:
                    paragraph.line_spacing = 0.9
                except Exception:
                    pass

            changed = True
        except Exception:
            pass

        # smaller
        for _ in range(4):
            overflow_ratio = self._estimate_overflow_ratio(shape, role)

            if overflow_ratio <= 1.03:
                break

            current_size = self._max_font_size(shape)
            min_size = self._min_font_size_for_role(role)

            if current_size <= min_size:
                break

            if overflow_ratio > 1.8:
                scale = 0.82
            elif overflow_ratio > 1.4:
                scale = 0.88
            else:
                scale = 0.93

            new_size = max(min_size, int(current_size * scale))

            if new_size >= current_size:
                break

            self._set_font_size(shape, new_size)
            changed = True

        # higher txt
        overflow_ratio = self._estimate_overflow_ratio(shape, role)

        if overflow_ratio > 1.08:
            expanded = self._expand_shape_height(
                shape=shape,
                slide_height=slide_height,
                role=role,
                overflow_ratio=overflow_ratio
            )

            if expanded:
                changed = True

        return changed

    def _estimate_overflow_ratio(self, shape, role: str) -> float:
        text = self._shape_text(shape)

        if not text:
            return 0.0

        try:
            width_in = int(shape.width) / EMU_PER_INCH
            height_in = int(shape.height) / EMU_PER_INCH
        except Exception:
            return 1.0

        if width_in <= 0 or height_in <= 0:
            return 1.0

        font_size = self._max_font_size(shape)

        if font_size <= 0:
            font_size = self._default_font_size_for_role(role)

        usable_width = max(0.2, width_in - 0.08)
        usable_height = max(0.15, height_in - 0.06)

        visual_capacity_per_line = max(
            5.0,
            usable_width * 12.0 / max(0.5, font_size / 12.0)
        )

        lines = []

        for raw_line in text.splitlines():
            raw_line = raw_line.strip()

            if raw_line:
                lines.append(raw_line)

        if not lines:
            lines = [text]

        required_lines = 0

        for line in lines:
            visual_width = self._visual_width(line)
            required_lines += max(1, math.ceil(visual_width / visual_capacity_per_line))

        line_height = (font_size / 72.0) * 1.18
        required_height = required_lines * line_height

        return required_height / max(0.01, usable_height)

    def _visual_width(self, text: str) -> float:
        width = 0.0

        for char in text:
            if char.isspace():
                width += 0.5
                continue

            east_asian_width = unicodedata.east_asian_width(char)

            if east_asian_width in {"F", "W"}:
                width += 2.0
            elif east_asian_width == "A":
                width += 1.5
            else:
                width += 1.0

        return width

    def _max_font_size(self, shape) -> int:
        sizes = []

        try:
            for paragraph in shape.text_frame.paragraphs:
                if paragraph.font.size:
                    sizes.append(paragraph.font.size.pt)

                for run in paragraph.runs:
                    if run.font.size:
                        sizes.append(run.font.size.pt)
        except Exception:
            pass

        if not sizes:
            role = self._guess_shape_role(shape)
            return self._default_font_size_for_role(role)

        return int(max(sizes))

    def _set_font_size(self, shape, font_size: int) -> None:
        try:
            for paragraph in shape.text_frame.paragraphs:
                try:
                    paragraph.font.size = Pt(font_size)
                except Exception:
                    pass

                for run in paragraph.runs:
                    try:
                        run.font.size = Pt(font_size)
                    except Exception:
                        pass
        except Exception:
            pass

    def _default_font_size_for_role(self, role: str) -> int:
        if role == "title":
            return 28

        if role == "subtitle":
            return 16

        if role == "label":
            return 12

        return 14

    def _min_font_size_for_role(self, role: str) -> int:
        if role == "title":
            return 20

        if role == "subtitle":
            return 13

        if role == "label":
            return 9

        return 10

    def _expand_shape_height(
        self,
        shape,
        slide_height: int,
        role: str,
        overflow_ratio: float
    ) -> bool:
        if role == "title":
            max_growth = Inches(0.25)
        elif role == "subtitle":
            max_growth = Inches(0.20)
        else:
            max_growth = Inches(0.55)

        try:
            current_height = int(shape.height)
            top = int(shape.top)
        except Exception:
            return False

        desired_growth = int(current_height * min(0.35, overflow_ratio - 1.0))
        desired_growth = max(0, min(desired_growth, int(max_growth)))

        if desired_growth <= 0:
            return False

        bottom_limit = slide_height - Inches(0.25)

        if top + current_height + desired_growth > bottom_limit:
            desired_growth = int(bottom_limit - top - current_height)

        if desired_growth <= 0:
            return False

        try:
            shape.height = current_height + desired_growth
            return True
        except Exception:
            return False

    def _resolve_overlaps(
        self,
        text_shapes: list[dict],
        slide_width: int,
        slide_height: int
    ) -> int:
        repaired = 0

        items = sorted(
            text_shapes,
            key=lambda item: (
                int(item["shape"].top),
                int(item["shape"].left)
            )
        )

        for i in range(len(items)):
            for j in range(i + 1, len(items)):
                a = items[i]
                b = items[j]

                if not self._boxes_overlap(a, b):
                    continue

                overlap_ratio = self._overlap_ratio(a, b)

                if overlap_ratio < 0.08:
                    continue

                moved = self._move_lower_shape_down(
                    upper=a,
                    lower=b,
                    slide_height=slide_height
                )

                if moved:
                    repaired += 1
                    continue


                target = b if b["role"] != "title" else a
                current_size = self._max_font_size(target["shape"])
                min_size = self._min_font_size_for_role(target["role"])

                if current_size > min_size:
                    self._set_font_size(target["shape"], max(min_size, current_size - 1))
                    repaired += 1

        return repaired

    def _boxes_overlap(self, a: dict, b: dict) -> bool:
        ax1, ay1, ax2, ay2 = self._box(a)
        bx1, by1, bx2, by2 = self._box(b)

        return ax1 < bx2 and ax2 > bx1 and ay1 < by2 and ay2 > by1

    def _overlap_ratio(self, a: dict, b: dict) -> float:
        ax1, ay1, ax2, ay2 = self._box(a)
        bx1, by1, bx2, by2 = self._box(b)

        ix1 = max(ax1, bx1)
        iy1 = max(ay1, by1)
        ix2 = min(ax2, bx2)
        iy2 = min(ay2, by2)

        if ix2 <= ix1 or iy2 <= iy1:
            return 0.0

        inter = (ix2 - ix1) * (iy2 - iy1)
        area_a = max(1, (ax2 - ax1) * (ay2 - ay1))
        area_b = max(1, (bx2 - bx1) * (by2 - by1))

        return inter / min(area_a, area_b)

    def _box(self, item: dict) -> tuple[int, int, int, int]:
        shape = item["shape"]

        left = int(shape.left)
        top = int(shape.top)
        width = int(shape.width)
        height = int(shape.height)

        return left, top, left + width, top + height

    def _move_lower_shape_down(
        self,
        upper: dict,
        lower: dict,
        slide_height: int
    ) -> bool:
        upper_shape = upper["shape"]
        lower_shape = lower["shape"]

        # dont move title
        if lower["role"] == "title":
            return False

        try:
            upper_bottom = int(upper_shape.top) + int(upper_shape.height)
            lower_top = int(lower_shape.top)
            lower_height = int(lower_shape.height)
        except Exception:
            return False

        padding = int(Inches(0.08))
        delta = upper_bottom - lower_top + padding

        if delta <= 0:
            return False

        max_delta = int(Inches(0.35))
        delta = min(delta, max_delta)

        bottom_limit = slide_height - Inches(0.25)

        if lower_top + delta + lower_height > bottom_limit:
            return False

        try:
            lower_shape.top = lower_top + delta
            return True
        except Exception:
            return False

    def _clamp_shape_to_slide(
        self,
        item: dict,
        slide_width: int,
        slide_height: int
    ) -> None:
        shape = item["shape"]

        try:
            left = int(shape.left)
            top = int(shape.top)
            width = int(shape.width)
            height = int(shape.height)
        except Exception:
            return

        margin = int(Inches(0.12))

        new_left = max(margin, min(left, slide_width - width - margin))
        new_top = max(margin, min(top, slide_height - height - margin))

        try:
            shape.left = new_left
            shape.top = new_top
        except Exception:
            pass
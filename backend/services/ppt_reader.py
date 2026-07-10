from pathlib import Path
from collections import Counter
import re

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


class PptReader:
    def extract_all_text(self, pptx_path: Path) -> list[str]:
        prs = Presentation(str(pptx_path))
        result = []

        for slide in prs.slides:
            lines = []

            for shape in self._iter_shapes(slide.shapes):
                if getattr(shape, "has_text_frame", False):
                    text = shape.text.strip()

                    if text:
                        lines.append(text)

            result.append("\n".join(lines))

        return result

    def extract_content_slides(self, pptx_path: Path) -> list[dict]:
        prs = Presentation(str(pptx_path))
        slides = []

        for index, slide in enumerate(prs.slides):
            clean_lines = []
            text_blocks = []
            image_count = 0
            table_count = 0

            for shape_index, shape in enumerate(self._iter_shapes(slide.shapes)):
                try:
                    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                        image_count += 1
                except Exception:
                    pass

                if getattr(shape, "has_table", False):
                    table_count += 1

                if getattr(shape, "has_text_frame", False):
                    text = shape.text.strip()

                    if not text:
                        continue

                    font_meta = self._get_font_meta(shape)

                    try:
                        x = float(shape.left / prs.slide_width)
                        y = float(shape.top / prs.slide_height)
                        w = float(shape.width / prs.slide_width)
                        h = float(shape.height / prs.slide_height)
                    except Exception:
                        x = 0.0
                        y = 0.0
                        w = 0.0
                        h = 0.0

                    block_lines = []

                    for line in text.splitlines():
                        line = line.strip()

                        if not line:
                            continue

                        if self._is_noise_line(line):
                            continue

                        block_lines.append(line)
                        clean_lines.append(line)

                    if block_lines:
                        text_blocks.append(
                            {
                                "shape_index": shape_index,
                                "text": "\n".join(block_lines),
                                "x": x,
                                "y": y,
                                "w": w,
                                "h": h,
                                "font_size": font_meta.get("font_size"),
                                "font_name": font_meta.get("font_name"),
                                "line_count": len(block_lines)
                            }
                        )

            title = self._guess_content_title(
                clean_lines=clean_lines,
                text_blocks=text_blocks,
                index=index
            )

            bullets = []

            for line in clean_lines:
                if line != title:
                    bullets.append(line)

            slides.append(
                {
                    "page_index": index,
                    "title": title,
                    "bullets": bullets,
                    "raw_text": "\n".join(clean_lines),
                    "text_blocks": text_blocks,
                    "image_count": image_count,
                    "table_count": table_count
                }
            )

        return slides

    def extract_reference_style(self, pptx_path: Path) -> dict:
        prs = Presentation(str(pptx_path))

        colors = []
        fonts = []
        font_sizes = []
        slide_summaries = []

        for slide_index, slide in enumerate(prs.slides):
            text_count = 0
            image_count = 0
            table_count = 0
            shape_count = len(slide.shapes)

            for shape in self._iter_shapes(slide.shapes):
                try:
                    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                        image_count += 1
                except Exception:
                    pass

                if getattr(shape, "has_table", False):
                    table_count += 1

                fill_color = self._shape_fill_color(shape)
                line_color = self._shape_line_color(shape)

                if fill_color:
                    colors.append(fill_color)

                if line_color:
                    colors.append(line_color)

                if getattr(shape, "has_text_frame", False):
                    text_count += 1

                    try:
                        for paragraph in shape.text_frame.paragraphs:
                            for run in paragraph.runs:
                                if run.font.name:
                                    fonts.append(run.font.name)

                                if run.font.size:
                                    font_sizes.append(float(run.font.size.pt))

                                try:
                                    if run.font.color and run.font.color.rgb:
                                        colors.append(self._rgb_to_hex(run.font.color.rgb))
                                except Exception:
                                    pass
                    except Exception:
                        pass

            slide_summaries.append(
                {
                    "page_index": slide_index,
                    "shape_count": shape_count,
                    "text_count": text_count,
                    "image_count": image_count,
                    "table_count": table_count
                }
            )

        color_list = self._most_common(colors, limit=12)
        font_list = self._most_common(fonts, limit=6)

        background_color = color_list[0] if color_list else "#F7F8FA"
        primary_text_color = color_list[1] if len(color_list) > 1 else "#A82A2A"
        secondary_text_color = color_list[2] if len(color_list) > 2 else "#8A1F1F"
        accent_color = color_list[3] if len(color_list) > 3 else "#A82A2A"

        return {
            "slide_width": prs.slide_width,
            "slide_height": prs.slide_height,
            "colors": color_list,
            "fonts": font_list,
            "font_sizes": font_sizes[:30],
            "background_color": background_color,
            "primary_text_color": primary_text_color,
            "secondary_text_color": secondary_text_color,
            "accent_color": accent_color,
            "title_font": font_list[0] if font_list else "Aptos Display",
            "body_font": font_list[1] if len(font_list) > 1 else "Aptos",
            "slide_summaries": slide_summaries
        }

    def extract_reference_templates(self, pptx_path: Path) -> list[dict]:
        prs = Presentation(str(pptx_path))
        templates = []

        for page_index, slide in enumerate(prs.slides):
            slots = []
            image_count = 0
            table_count = 0
            shape_count = len(slide.shapes)

            for shape in self._iter_shapes(slide.shapes):
                try:
                    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                        image_count += 1
                except Exception:
                    pass

                if getattr(shape, "has_table", False):
                    table_count += 1

                try:
                    x = float(shape.left / prs.slide_width)
                    y = float(shape.top / prs.slide_height)
                    w = float(shape.width / prs.slide_width)
                    h = float(shape.height / prs.slide_height)
                except Exception:
                    continue

                area = w * h

                if area <= 0:
                    continue

                if self._is_background_like(x, y, w, h, area):
                    continue

                has_text = bool(getattr(shape, "has_text_frame", False))
                fill_color = self._shape_fill_color(shape)
                line_color = self._shape_line_color(shape)

                font_meta = self._get_font_meta(shape) if has_text else {
                    "font_size": None,
                    "font_name": None
                }

                is_vertical_like = self._is_vertical_like_slot(
                    x=x,
                    y=y,
                    w=w,
                    h=h
                )

                role_guess = self._guess_slot_role(
                    x=x,
                    y=y,
                    w=w,
                    h=h,
                    font_size=font_meta.get("font_size"),
                    has_text=has_text,
                    is_vertical_like=is_vertical_like
                )

                usable_for_content = role_guess not in [
                    "footer",
                    "decorative_vertical",
                    "line",
                    "logo",
                    "tiny",
                    "background"
                ]

                if has_text:
                    slots.append(
                        {
                            "slot_index": len(slots),
                            "slot_kind": "text",
                            "role_guess": role_guess,
                            "usable_for_content": usable_for_content,
                            "x": x,
                            "y": y,
                            "w": w,
                            "h": h,
                            "area": area,
                            "font_size": font_meta.get("font_size"),
                            "font_name": font_meta.get("font_name"),
                            "fill_color": fill_color,
                            "line_color": line_color,
                            "is_dark": self._is_dark_color(fill_color),
                            "is_vertical_like": is_vertical_like
                        }
                    )

                elif self._looks_like_content_zone(
                    x=x,
                    y=y,
                    w=w,
                    h=h,
                    area=area,
                    fill_color=fill_color,
                    line_color=line_color
                ):
                    inset_x = 0.025
                    inset_y = 0.055

                    slots.append(
                        {
                            "slot_index": len(slots),
                            "slot_kind": "zone",
                            "role_guess": "card",
                            "usable_for_content": True,
                            "x": x + inset_x,
                            "y": y + inset_y,
                            "w": max(0.05, w - inset_x * 2),
                            "h": max(0.04, h - inset_y * 2),
                            "area": area,
                            "font_size": None,
                            "font_name": None,
                            "fill_color": fill_color,
                            "line_color": line_color,
                            "is_dark": self._is_dark_color(fill_color),
                            "is_vertical_like": False
                        }
                    )

            slots = self._deduplicate_slots(slots)
            slots = self._assign_slot_indexes(slots)

            templates.append(
                {
                    "page_index": page_index,
                    "shape_count": shape_count,
                    "image_count": image_count,
                    "table_count": table_count,
                    "slot_count": len(slots),
                    "slots": slots
                }
            )

        return templates

    def _iter_shapes(self, shapes):
        for shape in shapes:
            yield shape

            if hasattr(shape, "shapes"):
                try:
                    for child in self._iter_shapes(shape.shapes):
                        yield child
                except Exception:
                    pass

    def _is_noise_line(self, line: str) -> bool:
        clean = line.strip()

        if re.fullmatch(r"[•\-\s]*\d+\s*/\s*\d+", clean):
            return True

        if re.fullmatch(r"\d+\s*/\s*\d+", clean):
            return True

        if clean in ["•", "-", "—", "_"]:
            return True

        return False

    def _guess_content_title(
        self,
        clean_lines: list[str],
        text_blocks: list[dict],
        index: int
    ) -> str:
        if not clean_lines:
            return f"Slide {index + 1}"

        candidates = []

        for block in text_blocks:
            text = str(block.get("text", "")).strip()
            font_size = block.get("font_size") or 0
            y = block.get("y") or 0

            first_line = text.splitlines()[0].strip() if text else ""

            if first_line:
                candidates.append(
                    {
                        "text": first_line,
                        "font_size": font_size,
                        "y": y
                    }
                )

        if candidates:
            candidates = sorted(
                candidates,
                key=lambda item: (
                    item["font_size"],
                    -item["y"]
                ),
                reverse=True
            )

            return candidates[0]["text"]

        return clean_lines[0]

    def _is_background_like(
        self,
        x: float,
        y: float,
        w: float,
        h: float,
        area: float
    ) -> bool:
        if area > 0.70 and x < 0.08 and y < 0.08 and w > 0.80 and h > 0.80:
            return True

        return False

    def _looks_like_content_zone(
        self,
        x: float,
        y: float,
        w: float,
        h: float,
        area: float,
        fill_color: str | None,
        line_color: str | None
    ) -> bool:
        if area < 0.018:
            return False

        if w < 0.10 or h < 0.06:
            return False

        if h < 0.025 or w < 0.025:
            return False

        if w / max(h, 0.001) > 12:
            return False

        if h / max(w, 0.001) > 8:
            return False

        if y > 0.88 and area < 0.05:
            return False

        if fill_color or line_color:
            return True

        if area >= 0.04:
            return True

        return False

    def _is_vertical_like_slot(
        self,
        x: float,
        y: float,
        w: float,
        h: float
    ) -> bool:
        if w <= 0:
            return False

        return h / w >= 2.2 and w <= 0.16

    def _guess_slot_role(
        self,
        x: float,
        y: float,
        w: float,
        h: float,
        font_size,
        has_text: bool = True,
        is_vertical_like: bool = False
    ) -> str:
        font_size = font_size or 12
        area = w * h

        if area < 0.004:
            return "tiny"

        if is_vertical_like:
            return "decorative_vertical"

        if y > 0.86 and area < 0.05:
            return "footer"

        if h < 0.022 and w > 0.10:
            return "line"

        if y < 0.22 and font_size >= 22:
            return "title"

        if y < 0.36 and area < 0.18:
            return "subtitle"

        if w > 0.50 and h > 0.14:
            return "body"

        if area >= 0.035 and w < 0.55:
            return "card"

        if h < 0.10 and w > 0.18:
            return "label"

        return "content"

    def _deduplicate_slots(self, slots: list[dict]) -> list[dict]:
        result = []

        for slot in sorted(
            slots,
            key=lambda item: (
                float(item.get("y", 0)),
                float(item.get("x", 0)),
                -float(item.get("area", 0))
            )
        ):
            duplicate_index = None

            for i, existing in enumerate(result):
                overlap = self._slot_overlap_ratio(slot, existing)

                if overlap > 0.78:
                    duplicate_index = i
                    break

            if duplicate_index is None:
                result.append(slot)
                continue

            existing = result[duplicate_index]

            existing_area = float(existing.get("area", 0))
            slot_area = float(slot.get("area", 0))

            if slot.get("slot_kind") == "zone" and slot_area >= existing_area:
                result[duplicate_index] = slot

            elif existing.get("slot_kind") == "text" and slot.get("slot_kind") == "zone":
                if slot_area > existing_area * 1.5:
                    result[duplicate_index] = slot

            elif slot_area > existing_area * 1.8:
                result[duplicate_index] = slot

        return result

    def _assign_slot_indexes(self, slots: list[dict]) -> list[dict]:
        sorted_slots = sorted(
            slots,
            key=lambda item: (
                float(item.get("y", 0)),
                float(item.get("x", 0))
            )
        )

        for index, slot in enumerate(sorted_slots):
            slot["slot_index"] = index

        return sorted_slots

    def _slot_overlap_ratio(self, a: dict, b: dict) -> float:
        ax1 = a["x"]
        ay1 = a["y"]
        ax2 = a["x"] + a["w"]
        ay2 = a["y"] + a["h"]

        bx1 = b["x"]
        by1 = b["y"]
        bx2 = b["x"] + b["w"]
        by2 = b["y"] + b["h"]

        ix1 = max(ax1, bx1)
        iy1 = max(ay1, by1)
        ix2 = min(ax2, bx2)
        iy2 = min(ay2, by2)

        if ix2 <= ix1 or iy2 <= iy1:
            return 0.0

        inter = (ix2 - ix1) * (iy2 - iy1)
        area_a = a["w"] * a["h"]
        area_b = b["w"] * b["h"]

        return inter / max(0.0001, min(area_a, area_b))

    def _get_font_meta(self, shape) -> dict:
        font_sizes = []
        font_names = []

        try:
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.font.size:
                        font_sizes.append(float(run.font.size.pt))

                    if run.font.name:
                        font_names.append(run.font.name)
        except Exception:
            pass

        return {
            "font_size": max(font_sizes) if font_sizes else None,
            "font_name": font_names[0] if font_names else None
        }

    def _shape_fill_color(self, shape) -> str | None:
        try:
            fill = shape.fill

            if fill and fill.fore_color and fill.fore_color.rgb:
                return self._rgb_to_hex(fill.fore_color.rgb)
        except Exception:
            pass

        return None

    def _shape_line_color(self, shape) -> str | None:
        try:
            line = shape.line

            if line and line.color and line.color.rgb:
                return self._rgb_to_hex(line.color.rgb)
        except Exception:
            pass

        return None

    def _collect_shape_fill_color(self, shape, colors: list[str]):
        fill_color = self._shape_fill_color(shape)
        line_color = self._shape_line_color(shape)

        if fill_color:
            colors.append(fill_color)

        if line_color:
            colors.append(line_color)

    def _is_dark_color(self, hex_color: str | None) -> bool:
        if not hex_color:
            return False

        value = str(hex_color).replace("#", "").strip()

        if len(value) != 6:
            return False

        try:
            r = int(value[0:2], 16)
            g = int(value[2:4], 16)
            b = int(value[4:6], 16)

            brightness = 0.299 * r + 0.587 * g + 0.114 * b

            return brightness < 95
        except Exception:
            return False

    def _rgb_to_hex(self, rgb) -> str:
        value = str(rgb).strip()

        if value.startswith("#"):
            return value

        if len(value) == 6:
            return f"#{value}"

        return "#202431"

    def _most_common(self, items: list[str], limit: int = 10) -> list[str]:
        clean_items = []

        for item in items:
            if not item:
                continue

            value = str(item).strip()

            if value:
                clean_items.append(value)

        counter = Counter(clean_items)

        return [item for item, _ in counter.most_common(limit)]
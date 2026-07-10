from pathlib import Path
import time
import traceback
import math

from pptx import Presentation
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Pt
from pptx.dml.color import RGBColor


class PptWriter:
    def write_ai_design_plan(
        self,
        design_plan: dict,
        reference_style: dict,
        result_path: Path,
        reference_path: Path | None = None,
        reference_templates: list[dict] | None = None
    ):
        result_path = Path(result_path)
        result_path.parent.mkdir(parents=True, exist_ok=True)

        reference_templates = reference_templates or []

        if reference_path is not None and Path(reference_path).exists():
            try:
                self._copy_reference_slides_with_com(
                    design_plan=design_plan,
                    reference_path=Path(reference_path),
                    result_path=result_path
                )

                self._rewrite_copied_slides_content(
                    pptx_path=result_path,
                    design_plan=design_plan,
                    reference_style=reference_style,
                    reference_templates=reference_templates
                )

                return result_path

            except Exception:
                print("Template reuse with PowerPoint COM failed. Fallback to simple local renderer.")
                traceback.print_exc()

        self._write_simple_deck(
            design_plan=design_plan,
            reference_style=reference_style,
            result_path=result_path
        )

        return result_path

    def _copy_reference_slides_with_com(
        self,
        design_plan: dict,
        reference_path: Path,
        result_path: Path
    ):
        import comtypes
        import comtypes.client

        comtypes.CoInitialize()

        powerpoint = None
        source = None
        target = None

        try:
            powerpoint = comtypes.client.CreateObject("PowerPoint.Application")
            powerpoint.Visible = 1

            source = powerpoint.Presentations.Open(
                str(reference_path.resolve()),
                WithWindow=False
            )

            target = powerpoint.Presentations.Add()

            while target.Slides.Count > 0:
                target.Slides(1).Delete()

            reference_slide_count = source.Slides.Count
            output_slides = design_plan.get("slides", [])

            if reference_slide_count <= 0:
                raise RuntimeError("reference PPT has no slides")

            for output_index, slide_plan in enumerate(output_slides):
                selected = slide_plan.get(
                    "selected_reference_page_index",
                    output_index % reference_slide_count
                )

                try:
                    selected = int(selected)
                except Exception:
                    selected = output_index % reference_slide_count

                if selected < 0:
                    selected = 0

                if selected >= reference_slide_count:
                    selected = output_index % reference_slide_count

                source.Slides(selected + 1).Copy()
                time.sleep(0.1)

                pasted = target.Slides.Paste(target.Slides.Count + 1)
                time.sleep(0.1)

                try:
                    new_slide = pasted.Item(1)
                    self._clear_com_slide_text(new_slide)
                except Exception:
                    pass

            target.SaveAs(str(result_path.resolve()))

        finally:
            if source is not None:
                source.Close()

            if target is not None:
                target.Close()

            if powerpoint is not None:
                powerpoint.Quit()

            comtypes.CoUninitialize()

    def _clear_com_slide_text(self, slide):
        try:
            self._clear_com_shapes_text(slide.Shapes)
        except Exception:
            pass

    def _clear_com_shapes_text(self, shapes):
        for index in range(1, shapes.Count + 1):
            shape = shapes.Item(index)

            try:
                if shape.Type == 6:
                    self._clear_com_shapes_text(shape.GroupItems)
                    continue
            except Exception:
                pass

            try:
                if shape.HasTextFrame:
                    if shape.TextFrame.HasText:
                        shape.TextFrame.TextRange.Text = ""
            except Exception:
                pass

            try:
                if shape.HasTable:
                    table = shape.Table

                    for r in range(1, table.Rows.Count + 1):
                        for c in range(1, table.Columns.Count + 1):
                            table.Cell(r, c).Shape.TextFrame.TextRange.Text = ""
            except Exception:
                pass

    def _rewrite_copied_slides_content(
        self,
        pptx_path: Path,
        design_plan: dict,
        reference_style: dict,
        reference_templates: list[dict]
    ):
        prs = Presentation(str(pptx_path))
        plans = design_plan.get("slides", [])

        for slide_index, slide in enumerate(prs.slides):
            if slide_index >= len(plans):
                break

            plan = plans[slide_index]

            self._clear_all_template_text_with_python_pptx(slide)

            template = self._get_template_for_plan(
                plan=plan,
                reference_templates=reference_templates
            )

            raw_slots = template.get("slots", []) if template else []
            slots = self._filter_usable_slots(raw_slots)

            blocks = self._normalize_blocks_from_plan(plan)
            blocks = self._rebalance_blocks_for_slots(blocks, slots)

            print(
                f"Slide {slide_index + 1}: "
                f"{len(blocks)} content blocks, {len(slots)} usable slots"
            )

            self._place_blocks_by_template_slots(
                slide=slide,
                blocks=blocks,
                slots=slots,
                style=reference_style,
                slide_width=prs.slide_width,
                slide_height=prs.slide_height
            )

        prs.save(str(pptx_path))

    def _clear_all_template_text_with_python_pptx(self, slide):
        for shape in self._iter_shapes(slide.shapes):
            if not getattr(shape, "has_text_frame", False):
                continue

            try:
                shape.text_frame.clear()

                for paragraph in shape.text_frame.paragraphs:
                    paragraph.text = ""
            except Exception:
                try:
                    shape.text = ""
                except Exception:
                    pass

    def _iter_shapes(self, shapes):
        for shape in shapes:
            yield shape

            if hasattr(shape, "shapes"):
                try:
                    for child in self._iter_shapes(shape.shapes):
                        yield child
                except Exception:
                    pass

    def _get_template_for_plan(
        self,
        plan: dict,
        reference_templates: list[dict]
    ) -> dict | None:
        if not reference_templates:
            return None

        selected = plan.get("selected_reference_page_index", 0)

        try:
            selected = int(selected)
        except Exception:
            selected = 0

        if selected < 0 or selected >= len(reference_templates):
            selected = 0

        return reference_templates[selected]

    def _filter_usable_slots(self, slots: list[dict]) -> list[dict]:
        result = []

        for slot in slots:
            if not slot.get("usable_for_content", True):
                continue

            if slot.get("is_vertical_like"):
                continue

            role = slot.get("role_guess", "content")

            if role in ["footer", "decorative_vertical", "line", "logo", "tiny", "background"]:
                continue

            w = float(slot.get("w", 0))
            h = float(slot.get("h", 0))
            area = w * h

            if w < 0.08 or h < 0.04:
                continue

            if area < 0.006:
                continue

            result.append(slot)

        return sorted(
            result,
            key=lambda item: (
                float(item.get("y", 0)),
                float(item.get("x", 0)),
                -float(item.get("area", 0))
            )
        )

    def _normalize_blocks_from_plan(self, plan: dict) -> list[dict]:
        blocks = plan.get("content_blocks")

        if isinstance(blocks, list) and blocks:
            result = []

            for index, block in enumerate(blocks):
                if not isinstance(block, dict):
                    continue

                text = str(block.get("text", "")).strip()
                bullets = block.get("bullets", [])

                if isinstance(bullets, str):
                    bullets = [bullets]

                if not isinstance(bullets, list):
                    bullets = []

                bullets = [
                    str(item).strip()
                    for item in bullets
                    if item and str(item).strip()
                ]

                if not text and not bullets:
                    continue

                result.append(
                    {
                        "block_id": str(block.get("block_id", f"block_{index + 1}")),
                        "role": self._clean_role(block.get("role", "body")),
                        "target_slot_index": self._safe_int_or_none(
                            block.get("target_slot_index")
                        ),
                        "text": text,
                        "bullets": bullets
                    }
                )

            if result:
                return result

        return self._fallback_blocks(plan)

    def _fallback_blocks(self, plan: dict) -> list[dict]:
        result = []

        title = str(plan.get("title", "")).strip()
        subtitle = str(plan.get("subtitle", "")).strip()
        bullets = plan.get("bullets", [])

        if not isinstance(bullets, list):
            bullets = []

        if title:
            result.append(
                {
                    "block_id": "title_1",
                    "role": "title",
                    "target_slot_index": None,
                    "text": title,
                    "bullets": []
                }
            )

        if subtitle:
            result.append(
                {
                    "block_id": "subtitle_1",
                    "role": "subtitle",
                    "target_slot_index": None,
                    "text": subtitle,
                    "bullets": []
                }
            )

        for i, group in enumerate(self._split_bullets(bullets)):
            result.append(
                {
                    "block_id": f"body_{i + 1}",
                    "role": "body",
                    "target_slot_index": None,
                    "text": "",
                    "bullets": group
                }
            )

        return result

    def _rebalance_blocks_for_slots(
        self,
        blocks: list[dict],
        slots: list[dict]
    ) -> list[dict]:
        content_slots = [
            slot for slot in slots
            if slot.get("role_guess") in ["card", "content", "body", "label"]
        ]

        if len(content_slots) < 2:
            return blocks

        result = []

        for block in blocks:
            role = block.get("role", "body")
            bullets = block.get("bullets", [])
            text = block.get("text", "")

            if role in ["title", "subtitle", "kpi", "quote"]:
                result.append(block)
                continue

            if not isinstance(bullets, list):
                result.append(block)
                continue

            if len(bullets) <= 1:
                result.append(block)
                continue

            max_groups = min(len(content_slots), len(bullets), 4)
            chunk_size = math.ceil(len(bullets) / max_groups)

            for i in range(0, len(bullets), chunk_size):
                group = bullets[i:i + chunk_size]

                result.append(
                    {
                        "block_id": f"{block.get('block_id', 'body')}_{i // chunk_size + 1}",
                        "role": "card" if len(content_slots) >= 2 else role,
                        "target_slot_index": None,
                        "text": text if i == 0 else "",
                        "bullets": group
                    }
                )

        return result

    def _place_blocks_by_template_slots(
        self,
        slide,
        blocks: list[dict],
        slots: list[dict],
        style: dict,
        slide_width,
        slide_height
    ):
        used_slot_indexes = set()
        used_boxes = []

        for block_index, block in enumerate(blocks):
            slot = self._choose_slot(
                block=block,
                slots=slots,
                used_slot_indexes=used_slot_indexes,
                used_boxes=used_boxes
            )

            if slot is not None:
                used_slot_indexes.add(slot["slot_index"])
                box = self._slot_to_box(slot, slide_width, slide_height)
            else:
                box = self._fallback_box(
                    block=block,
                    block_index=block_index,
                    slide_width=slide_width,
                    slide_height=slide_height
                )

            used_boxes.append(box)

            self._add_horizontal_textbox(
                slide=slide,
                box=box,
                block=block,
                slot=slot,
                style=style
            )

    def _choose_slot(
        self,
        block: dict,
        slots: list[dict],
        used_slot_indexes: set[int],
        used_boxes: list[dict]
    ) -> dict | None:
        role = block.get("role", "body")
        target = block.get("target_slot_index")

        if target is not None:
            for slot in slots:
                slot_index = slot.get("slot_index")

                if slot_index != target:
                    continue

                if slot_index in used_slot_indexes:
                    continue

                if self._slot_bad_for_role(slot, role):
                    continue

                if self._slot_overlaps_used(slot, used_boxes):
                    continue

                return slot

        priority = {
            "title": ["title", "body", "content"],
            "subtitle": ["subtitle", "label", "body", "content", "card"],
            "body": ["body", "content", "card"],
            "card": ["card", "content", "body"],
            "step": ["card", "content", "body"],
            "kpi": ["card", "content"],
            "quote": ["body", "content", "card"],
            "note": ["label", "content", "card"],
            "footer": ["content"]
        }

        wanted_roles = priority.get(role, ["body", "content", "card"])

        candidates = []

        for wanted in wanted_roles:
            for slot in slots:
                slot_index = slot.get("slot_index")

                if slot_index in used_slot_indexes:
                    continue

                if slot.get("role_guess") != wanted:
                    continue

                if self._slot_bad_for_role(slot, role):
                    continue

                if self._slot_overlaps_used(slot, used_boxes):
                    continue

                candidates.append(slot)

            if candidates:
                return self._best_slot_for_role(candidates, role)

        for slot in slots:
            slot_index = slot.get("slot_index")

            if slot_index in used_slot_indexes:
                continue

            if self._slot_bad_for_role(slot, role):
                continue

            if self._slot_overlaps_used(slot, used_boxes):
                continue

            return slot

        return None

    def _best_slot_for_role(self, candidates: list[dict], role: str) -> dict:
        if role == "title":
            return sorted(
                candidates,
                key=lambda s: (
                    float(s.get("y", 0)),
                    -float(s.get("area", 0))
                )
            )[0]

        if role == "subtitle":
            return sorted(
                candidates,
                key=lambda s: (
                    float(s.get("y", 0)),
                    float(s.get("x", 0))
                )
            )[0]

        return sorted(
            candidates,
            key=lambda s: (
                float(s.get("y", 0)),
                float(s.get("x", 0)),
                -float(s.get("area", 0))
            )
        )[0]

    def _slot_bad_for_role(self, slot: dict, role: str) -> bool:
        if slot.get("is_vertical_like"):
            return True

        slot_role = slot.get("role_guess", "content")

        if slot_role in ["footer", "decorative_vertical", "line", "logo", "tiny", "background"]:
            return True

        w = float(slot.get("w", 0))
        h = float(slot.get("h", 0))

        if role in ["body", "card", "step", "quote"] and h < 0.07:
            return True

        if role == "title" and w < 0.18:
            return True

        return False

    def _slot_overlaps_used(
        self,
        slot: dict,
        used_boxes: list[dict]
    ) -> bool:
        if not used_boxes:
            return False

        slot_box = {
            "left": float(slot.get("x", 0)),
            "top": float(slot.get("y", 0)),
            "width": float(slot.get("w", 0)),
            "height": float(slot.get("h", 0))
        }

        for box in used_boxes:
            normalized_box = {
                "left": box.get("norm_x", 0),
                "top": box.get("norm_y", 0),
                "width": box.get("norm_w", 0),
                "height": box.get("norm_h", 0)
            }

            if self._box_overlap_ratio(slot_box, normalized_box) > 0.45:
                return True

        return False

    def _box_overlap_ratio(self, a: dict, b: dict) -> float:
        ax1 = a["left"]
        ay1 = a["top"]
        ax2 = a["left"] + a["width"]
        ay2 = a["top"] + a["height"]

        bx1 = b["left"]
        by1 = b["top"]
        bx2 = b["left"] + b["width"]
        by2 = b["top"] + b["height"]

        ix1 = max(ax1, bx1)
        iy1 = max(ay1, by1)
        ix2 = min(ax2, bx2)
        iy2 = min(ay2, by2)

        if ix2 <= ix1 or iy2 <= iy1:
            return 0.0

        inter = (ix2 - ix1) * (iy2 - iy1)
        area_a = a["width"] * a["height"]
        area_b = b["width"] * b["height"]

        return inter / max(0.0001, min(area_a, area_b))

    def _slot_to_box(self, slot: dict, slide_width, slide_height) -> dict:
        x = float(slot.get("x", 0.08))
        y = float(slot.get("y", 0.18))
        w = float(slot.get("w", 0.50))
        h = float(slot.get("h", 0.20))

        return {
            "left": int(slide_width * x),
            "top": int(slide_height * y),
            "width": int(slide_width * w),
            "height": int(slide_height * h),
            "norm_x": x,
            "norm_y": y,
            "norm_w": w,
            "norm_h": h
        }

    def _fallback_box(
        self,
        block: dict,
        block_index: int,
        slide_width,
        slide_height
    ) -> dict:
        role = block.get("role", "body")

        if role == "title":
            x = 0.08
            y = 0.08
            w = 0.70
            h = 0.12
        elif role == "subtitle":
            x = 0.10
            y = 0.22
            w = 0.65
            h = 0.10
        else:
            content_index = max(0, block_index - 2)
            col = content_index % 2
            row = content_index // 2

            x = 0.10 + col * 0.42
            y = 0.42 + row * 0.20
            w = 0.36
            h = 0.16

        return {
            "left": int(slide_width * x),
            "top": int(slide_height * y),
            "width": int(slide_width * w),
            "height": int(slide_height * h),
            "norm_x": x,
            "norm_y": y,
            "norm_w": w,
            "norm_h": h
        }

    def _add_horizontal_textbox(
        self,
        slide,
        box: dict,
        block: dict,
        slot: dict | None,
        style: dict
    ):
        textbox = slide.shapes.add_textbox(
            box["left"],
            box["top"],
            box["width"],
            box["height"]
        )

        frame = textbox.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.vertical_anchor = MSO_ANCHOR.TOP
        frame.margin_left = 0
        frame.margin_right = 0
        frame.margin_top = 0
        frame.margin_bottom = 0

        role = block.get("role", "body")
        text = str(block.get("text", "")).strip()
        bullets = block.get("bullets", [])

        if not isinstance(bullets, list):
            bullets = []

        font_size = self._font_size_for_block(role, text, bullets, box)
        color = self._color_for_block(role, slot, style)
        font_name = self._title_font(style) if role == "title" else self._body_font(style)

        first = frame.paragraphs[0]

        if text:
            first.text = text
            first.font.name = font_name
            first.font.size = Pt(font_size)
            first.font.bold = role in ["title", "kpi"]
            first.font.color.rgb = self._rgb(color)
            first.alignment = PP_ALIGN.LEFT
            first.space_after = Pt(4)
        else:
            first.text = ""

        for bullet in bullets[:6]:
            p = frame.add_paragraph()
            p.text = f"• {str(bullet).strip()}"
            p.font.name = self._body_font(style)
            p.font.size = Pt(max(8, font_size - 2))
            p.font.bold = False
            p.font.color.rgb = self._rgb(color)
            p.space_after = Pt(2)
            p.level = 0

    def _font_size_for_block(
        self,
        role: str,
        text: str,
        bullets: list[str],
        box: dict
    ) -> int:
        width = box["width"]
        height = box["height"]

        short_box = height < 750000
        narrow_box = width < 2800000

        total_chars = len(text) + sum(len(str(b)) for b in bullets)

        if role == "title":
            if total_chars <= 10:
                return 34
            if total_chars <= 20:
                return 28
            return 22

        if role == "subtitle":
            return 15 if total_chars < 60 else 12

        if role == "kpi":
            return 24

        if narrow_box or short_box:
            if total_chars > 90:
                return 9
            if total_chars > 50:
                return 10
            return 12

        if total_chars > 180:
            return 10

        if total_chars > 100:
            return 12

        return 14

    def _write_simple_deck(
        self,
        design_plan: dict,
        reference_style: dict,
        result_path: Path
    ):
        prs = Presentation()
        prs.slide_width = reference_style.get("slide_width", prs.slide_width)
        prs.slide_height = reference_style.get("slide_height", prs.slide_height)

        blank_layout = prs.slide_layouts[6]

        for plan in design_plan.get("slides", []):
            slide = prs.slides.add_slide(blank_layout)

            fill = slide.background.fill
            fill.solid()
            fill.fore_color.rgb = self._rgb(self._background(reference_style))

            blocks = self._normalize_blocks_from_plan(plan)

            self._place_blocks_by_template_slots(
                slide=slide,
                blocks=blocks,
                slots=[],
                style=reference_style,
                slide_width=prs.slide_width,
                slide_height=prs.slide_height
            )

        prs.save(str(result_path))

    def _split_bullets(self, bullets: list[str]) -> list[list[str]]:
        clean = [
            str(item).strip()
            for item in bullets
            if item and str(item).strip()
        ]

        if not clean:
            return []

        if len(clean) <= 3:
            return [clean]

        result = []

        for i in range(0, len(clean), 3):
            result.append(clean[i:i + 3])

        return result[:4]

    def _clean_role(self, role) -> str:
        value = str(role).strip().lower()

        allowed = {
            "title",
            "subtitle",
            "body",
            "card",
            "step",
            "kpi",
            "quote",
            "note",
            "footer"
        }

        if value in allowed:
            return value

        return "body"

    def _safe_int_or_none(self, value):
        try:
            if value is None:
                return None

            return int(value)
        except Exception:
            return None

    def _color_for_block(
        self,
        role: str,
        slot: dict | None,
        style: dict
    ) -> str:
        if slot and slot.get("is_dark"):
            return "#FFFFFF"

        if role in ["subtitle", "note", "footer"]:
            return self._secondary(style)

        return self._primary(style)

    def _rgb(self, hex_color: str) -> RGBColor:
        if not hex_color:
            hex_color = "#A82A2A"

        value = str(hex_color).replace("#", "").strip()

        if len(value) != 6:
            value = "A82A2A"

        try:
            return RGBColor(
                int(value[0:2], 16),
                int(value[2:4], 16),
                int(value[4:6], 16)
            )
        except Exception:
            return RGBColor(168, 42, 42)

    def _background(self, style: dict) -> str:
        return style.get("background_color", "#F7F8FA")

    def _primary(self, style: dict) -> str:
        return style.get("primary_text_color", "#A82A2A")

    def _secondary(self, style: dict) -> str:
        return style.get("secondary_text_color", "#8A1F1F")

    def _title_font(self, style: dict) -> str:
        return style.get("title_font", "Aptos Display")

    def _body_font(self, style: dict) -> str:
        return style.get("body_font", "Aptos")